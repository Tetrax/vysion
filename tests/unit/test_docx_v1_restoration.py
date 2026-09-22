import re
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

import pytest
from docx import Document

from vysion.adapters.fortiguard import FortiGuardResult, FortiGuardStatus
from vysion.audit.engine import AuditEngine
from vysion.audit.models import AuditContext, WanSelection, WanSelectionKind
from vysion.audit.parser import FortiGateParser
from vysion.audit.registry import default_registry
from vysion.reports.docx_report import render_docx
from vysion.reports.json_report import JsonAuditReport
from vysion.reports.presentation import build_presentation, present_findings

CONFIGURATION = Path("/home/hermes/.hermes/attachments/FW-AVR-01_7-2_1639_202608121044.conf")
SECOND_CONFIGURATION = Path(
    "/home/hermes/.hermes/attachments/CF-UK-WIGAN-MASTER_7-4_2829_202609090946.conf"
)

pytestmark = pytest.mark.skipif(
    not CONFIGURATION.exists(),
    reason=(
        "real FortiGate backup required: place the field backup under "
        "/home/hermes/.hermes/attachments/ (see OPERATIONS.md)"
    ),
)


def _report_for(configuration_path: Path, context: AuditContext | None = None) -> JsonAuditReport:
    context = context or AuditContext()
    configuration = FortiGateParser().parse(configuration_path.read_text(encoding="utf-8"))
    findings = tuple(AuditEngine(default_registry()).run(configuration, context))
    presented = present_findings(findings)
    created_at = datetime(2026, 8, 20, 12, 0, tzinfo=UTC)
    return JsonAuditReport(
        report_id=uuid4(),
        created_at=created_at,
        expires_at=created_at + timedelta(minutes=5),
        source_name=configuration_path.name,
        context=context,
        device_identity=configuration.device_identity,
        fortiguard=FortiGuardResult(
            status=FortiGuardStatus.UNKNOWN,
            detail="fixture",
        ),
        findings=presented,
        presentation=build_presentation(findings, configuration, context),
    )


def _real_report() -> JsonAuditReport:
    return _report_for(CONFIGURATION)


@pytest.mark.skipif(
    not SECOND_CONFIGURATION.exists(),
    reason="second field backup not available on this machine",
)
def test_cf_uk_wigan_utm_policy_without_logtraffic_is_unknown() -> None:
    configuration = FortiGateParser().parse(SECOND_CONFIGURATION.read_text(encoding="utf-8"))
    context = AuditContext(
        wan_selections=(
            WanSelection(
                name="wan1",
                kind=WanSelectionKind.INTERFACE,
                interfaces=("wan1",),
            ),
        )
    )

    finding = next(
        item
        for item in AuditEngine(default_registry()).run(configuration, context)
        if item.control_id == "FW-UTM-PROFILE-BINDING-001"
    )

    assert finding.status.value == "UNKNOWN"
    assert finding.applicability.value == "unknown"
    assert finding.affected_objects


@pytest.mark.skipif(
    not SECOND_CONFIGURATION.exists(),
    reason="second field backup not available on this machine",
)
def test_cf_uk_wigan_utm_unknown_docx_uses_business_wording() -> None:
    context = AuditContext(
        wan_selections=(
            WanSelection(
                name="wan1",
                kind=WanSelectionKind.INTERFACE,
                interfaces=("wan1",),
            ),
        ),
        utm_license=True,
    )

    body = _client_body(Document(BytesIO(render_docx(_report_for(SECOND_CONFIGURATION, context)))))

    assert (
        "La configuration disponible ne permet pas d'établir le comportement de "
        "journalisation de toutes les règles avec inspection UTM activée."
    ) in body
    assert "logtraffic" not in body.casefold()


def test_docx_uses_the_v1_document_order_and_keeps_ssl_vpn_na() -> None:
    document = Document(BytesIO(render_docx(_real_report())))
    headings = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading")
    ]

    expected = [
        "3. Audit de configuration",
        "3.1 Audit - Système",
        "3.1.1 Version Fortigate",
        "3.1.2 Modèle Fortigate",
        "3.1.3 Sauvegardes automatiques",
        "3.1.4 Objets sans référence",
        "3.1.5 Auto-installation d'image par USB",
        "3.2 Audit - Administration et comptes",
        "3.2.1 Port HTTPS personnalisé",
        "3.2.2 Synchronisation avec un FortiManager",
        "3.2.3 Synchronisation avec un FortiAnalyzer",
        "3.2.4 Durcissement accès administration à votre Fortigate",
        "3.2.5 Compte 'Admin' par défaut",
        "3.2.6 MFA pour les comptes admins et utilisateurs",
        "3.3 Audit - Réseaux et flux",
        "3.3.1 Règles en 'By Sequence'",
        "3.3.2 Logs sur la règle implicit deny",
        "3.3.3 Utilisation du SD-WAN",
        "3.3.4 Blocage des ISDB malveillants",
        "3.3.5 Filtrage des ports vers Internet",
        "3.3.6 Absence de VIP en ANY",
        "3.3.7 Absence de Virtual Server en ANY",
        "3.3.8 Utilisation de la GEO-IP",
        "3.3.9 Utilisation de nos CTI",
        "3.3.10 Logs en UTM sans profil de sécurité",
        "3.3.11 Route Blackhole pour les réseaux privés",
        "3.3.12 Ports-Deny vers Internet",
        "3.3.13 Utilisation du LDAPS",
        "3.4 Audit - Connexions distantes : VPN",
        "3.4.1 Utilisation du VPN SSL",
        "3.4.2 Durcissement des VPN IPSEC : contrôle IKE",
        "3.4.3 Durcissement des VPN IPSEC : contrôle DH Group",
        "3.4.4 Durcissement des VPN IPSEC : ESP et algorithmes",
        "3.5 Audit - Profils de sécurité UTM",
        "3.5.1 Vérification de la licence UTM",
        "3.5.2 Mises à jour FortiGuard",
        "3.5.3 FortiSandbox Cloud",
        "3.5.4 Utilisation du DNS-Filter",
        "3.5.5 Utilisation du Web-Filter",
        "3.5.6 Utilisation de l'Antivirus",
        "3.5.7 Utilisation de l'IPS",
        "3.5.8 Utilisation de l'Application Control",
    ]
    positions = [headings.index(value) for value in expected]
    assert positions == sorted(positions)
    assert "NON APPLICABLE" in "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_docx_keeps_the_v1_logo_placeholder_without_client_logo() -> None:
    document = Document(BytesIO(render_docx(_real_report())))
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "LOGO CLIENT si existe" in body


def test_docx_inserts_a_client_logo_when_one_is_available() -> None:
    logo = Path(__file__).parents[2] / "src" / "vysion" / "reports" / "assets" / "cluster.png"
    document = Document(BytesIO(render_docx(_real_report(), client_logo_path=logo)))
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "LOGO CLIENT si existe" not in body
    # Client logo plus the three business diagrams of the mandatory body
    # (ISDB, CTI, security profiles); the Cluster diagram left with the
    # removed Cluster section.
    assert len(document.inline_shapes) >= 4


def test_docx_embeds_v1_business_diagrams_in_the_report_body() -> None:
    with ZipFile(BytesIO(render_docx(_real_report()))) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")

    # The mandatory body carries three business diagrams (ISDB, CTI, security
    # profiles); the Cluster diagram left with the removed Cluster section.
    assert document_xml.count("<pic:pic") >= 3


# Internal migration/architecture vocabulary that must never reach the client
# document.  Patterns stay qualified (V1/V2 next to architecture words, or
# explicit internal phrases) so a legitimate client value like an object name
# containing "V2" is not flagged.
INTERNAL_VOCABULARY_PATTERNS = (
    r"périmètre\s+V2",
    r"contrôle historique",
    r"règle historique",
    r"différence de couverture",
    r"\bV1\s*/\s*V2\b",
    r"\bV2-only\b",
    r"\blegacy_v1\b",
    r"\b(?:moteur|architecture|parité|baseline|gate)\s+V[12]\b",
)


def test_docx_client_body_has_no_internal_architecture_vocabulary() -> None:
    document = Document(BytesIO(render_docx(_real_report())))
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in row.cells)
    body = "\n".join(texts)

    for pattern in INTERNAL_VOCABULARY_PATTERNS:
        assert re.search(pattern, body, flags=re.IGNORECASE) is None, pattern
    assert "adresses, groupes d'adresses, VIP, groupes de VIP, Virtual Server, zones" in body


def test_docx_vpn_ssl_na_uses_the_v1_client_wording() -> None:
    document = Document(BytesIO(render_docx(_real_report())))
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "Le VPN SSL n'est pas utilisé." in body


# FortiOS implementation vocabulary observed in the field validation report
# (CFG-UNUSED wording, implicit-deny logging, UTM profile families, GEO-IP
# selection, LDAPS messages, SD-WAN selection).  Patterns stay qualified so a
# legitimate client object name is not flagged: exact identifiers or English
# multi-word phrases, never a bare French word.
IMPLEMENTATION_VOCABULARY_PATTERNS = (
    r"\bfwpolicy-implicit-log\b",
    r"\blog setting\b",
    r"\b(?:enable|disable)\b",
    r"\busage prouvé\b",
    r"\btyped\b",
    r"\bantivirus profile\b",
    r"\bips sensor\b",
    r"\bapplication list\b",
    r"\bwebfilter profile\b",
    r"\bdnsfilter profile\b",
    r"\bbackup complet\b",
    r"\buser ldap\b",
    r"\brouter static\b",
    r"\bvpn ssl settings\b",
    r"\bnamespace",
    r"\bfortigate-[a-z-]+",
    r"défaut automatique V[12]",
    r"explicitement non sécurisé",
    r"\(\s*\)",
)


def _client_body(document) -> str:
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in row.cells)
    return "\n".join(texts)


@pytest.mark.parametrize(
    "configuration_path",
    [
        CONFIGURATION,
        pytest.param(
            SECOND_CONFIGURATION,
            marks=pytest.mark.skipif(
                not SECOND_CONFIGURATION.exists(),
                reason="second field backup not available on this machine",
            ),
        ),
    ],
    ids=["FW-AVR-01", "CF-UK-WIGAN-MASTER"],
)
def test_docx_client_body_has_no_fortios_implementation_vocabulary(
    configuration_path: Path,
) -> None:
    document = Document(BytesIO(render_docx(_report_for(configuration_path))))
    body = _client_body(document)

    for pattern in IMPLEMENTATION_VOCABULARY_PATTERNS:
        match = re.search(pattern, body, flags=re.IGNORECASE)
        assert match is None, f"{pattern!r} matched {match.group(0)!r}"
