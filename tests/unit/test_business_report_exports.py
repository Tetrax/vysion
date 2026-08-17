from datetime import UTC, datetime, timedelta
from io import BytesIO
from uuid import uuid4

from docx import Document
from openpyxl import load_workbook

from vysion.adapters.fortiguard import FortiGuardResult, FortiGuardStatus
from vysion.audit.models import (
    AffectedObject,
    Applicability,
    AuditContext,
    AuditFinding,
    AuditPriority,
    AuditSeverity,
    AuditStatus,
    EvidenceItem,
    RiskAssessment,
    RuleMatchStatistics,
)
from vysion.reports.docx_report import render_docx
from vysion.reports.json_report import AccountMetadata, EquipmentMetadata, JsonAuditReport
from vysion.reports.xlsx_report import render_xlsx


def business_report() -> JsonAuditReport:
    created_at = datetime(2026, 8, 17, 12, 0, tzinfo=UTC)
    return JsonAuditReport(
        report_id=uuid4(),
        created_at=created_at,
        expires_at=created_at + timedelta(minutes=5),
        source_name="client.conf",
        context=AuditContext(
            client="Client métier",
            site="Paris-DC1",
            serial_number="FG100",
            uptime="42 days",
            operator_comment="HA à confirmer avec l’exploitant.",
            rule_match_statistics=RuleMatchStatistics(unmatched_rules=7),
            operator="Orange Business",
            ha=True,
            mpls=False,
            utm_license_details={
                "status": "active",
                "expiration_date": "2027-03-31",
                "provenance": "FortiManager",
                "manual": False,
            },
        ),
        equipment=EquipmentMetadata(
            hostname="client-fw",
            model="100F",
            firmware_version="7.4.3",
            serial_number="FG100",
            interface_names=("wan1", "lan1"),
            zone_names=("internet",),
            sdwan_zone_names=("virtual-wan-link",),
            interface_zone_relations=("wan1 → internet",),
        ),
        accounts=(
            AccountMetadata(name="admin", kind="administrator", two_factor="enable"),
            AccountMetadata(name="svc-backup", kind="local-user", two_factor="disable"),
        ),
        fortiguard=FortiGuardResult(status=FortiGuardStatus.AVAILABLE, detail="ruleset"),
        findings=(
            AuditFinding(
                control_id="NET-WAN-MGMT-001",
                title="Administration exposée",
                status=AuditStatus.FAIL,
                category="network",
                priority=AuditPriority.P0,
                severity=AuditSeverity.HIGH,
                applicability=Applicability.APPLICABLE,
                evidence_items=(
                    EvidenceItem(
                        section="system interface",
                        entry="wan1",
                        directive="allowaccess",
                        tokens=("https",),
                        line=10,
                    ),
                ),
                affected_objects=(AffectedObject(name="wan1", object_type="interface"),),
                message="Une administration est exposée sur le WAN.",
                risk=RiskAssessment(
                    summary="Compromission du plan de gestion.",
                    impact="Élevé",
                    likelihood="Vraisemblable",
                    treatment="Simple",
                ),
                recommendation="Retirer HTTPS de allowaccess.",
                remediation="Modifier l'interface puis valider la configuration.",
                customer_approval=False,
            ),
        ),
    )


def test_json_report_contains_equipment_and_account_metadata_for_all_renderers() -> None:
    report = business_report()

    payload = report.model_dump(mode="json")

    assert payload["equipment"]["hostname"] == "client-fw"
    assert payload["equipment"]["interface_zone_relations"] == ["wan1 → internet"]
    assert payload["accounts"][0]["name"] == "admin"


def test_docx_has_client_cover_context_summary_and_separate_technical_detail() -> None:
    report = business_report()

    document = Document(BytesIO(render_docx(report)))
    text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )

    for heading in (
        "Rapport d’audit Vysion",
        "Confidentialité",
        "Contexte",
        "Échelle de risque",
        "Synthèse",
        "Détail des contrôles",
        "Tableau récapitulatif final des risques",
        "Vue client",
        "Vue technique",
        "Réseau",
    ):
        assert heading in text
    assert "Client métier" in text
    assert "allowaccess" in text
    assert "wan1 → internet" in text
    assert "Règles sans match" in text
    assert "Hit Count <= 0" in text
    assert "HA à confirmer avec l’exploitant." in text


def test_xlsx_has_business_tabs_and_keeps_formula_injection_protection() -> None:
    workbook = load_workbook(BytesIO(render_xlsx(business_report())), data_only=False)

    required = {
        "Audit configuration",
        "Actions sans accord",
        "Actions avec accord",
        "Statistiques",
        "Comptes",
        "Métadonnées équipement",
    }
    assert required <= set(workbook.sheetnames)
    assert workbook["Audit configuration"]["A1"].value == "ID"
    assert workbook["Audit configuration"]["D1"].value == "Résultat"
    assert workbook["Audit configuration"]["A2"].value == "NET-WAN-MGMT-001"
    assert workbook["Actions sans accord"]["A2"].value == "NET-WAN-MGMT-001"
    assert workbook["Actions avec accord"].max_row == 1
    assert workbook["Comptes"]["A2"].value == "admin"
    assert workbook["Métadonnées équipement"]["B1"].value == "client-fw"
    assert workbook["Statistiques"]["B2"].value == 1
    assert workbook["Synthèse"]["B7"].alignment.wrap_text is True
    assert workbook["Audit configuration"]["F2"].alignment.wrap_text is True
    assert any(
        row[0].value == "Commentaire contexte"
        and row[1].value == "HA à confirmer avec l’exploitant."
        for row in workbook["Synthèse"].iter_rows(min_row=1, max_col=2)
    )
    statistics = {
        str(row[0].value): row[1].value
        for row in workbook["Statistiques"].iter_rows(min_row=1, max_col=2)
        if row[0].value is not None
    }
    assert statistics["Règles sans match (observation opérateur)"] == 7
    assert statistics["Méthode règles sans match"] == "Firewall policy Hit Count <= 0"


def test_xlsx_statistics_escape_operator_controlled_strings() -> None:
    report = business_report()
    context = report.context.model_copy(
        update={
            "rule_match_statistics": RuleMatchStatistics(
                unmatched_rules=7,
                source='=HYPERLINK("https://evil.invalid")',
                method="@external-command",
            )
        }
    )
    workbook = load_workbook(
        BytesIO(render_xlsx(report.model_copy(update={"context": context}))),
        data_only=False,
    )

    statistics = {
        str(row[0].value): row[1]
        for row in workbook["Statistiques"].iter_rows(min_row=1, max_col=2)
        if row[0].value is not None
    }
    assert statistics["Source règles sans match"].value == "'=HYPERLINK(\"https://evil.invalid\")"
    assert statistics["Source règles sans match"].data_type == "s"
    assert statistics["Méthode règles sans match"].value == "'@external-command"
    assert statistics["Méthode règles sans match"].data_type == "s"
