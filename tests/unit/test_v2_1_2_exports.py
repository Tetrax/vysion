from datetime import UTC, datetime, timedelta
from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile

from openpyxl import load_workbook

from vysion.adapters.fortiguard import FortiGuardResult, FortiGuardStatus
from vysion.audit.models import Applicability, AuditFinding, AuditStatus
from vysion.reports.docx_report import render_docx
from vysion.reports.json_report import JsonAuditReport
from vysion.reports.xlsx_report import render_xlsx


def test_not_applicable_status_is_preserved_in_machine_exports_but_hidden_from_docx() -> None:
    created_at = datetime(2026, 8, 14, 12, 0, tzinfo=UTC)
    report = JsonAuditReport(
        report_id=uuid4(),
        created_at=created_at,
        expires_at=created_at + timedelta(minutes=5),
        source_name="synthetic.conf",
        fortiguard=FortiGuardResult(status=FortiGuardStatus.UNKNOWN, detail="synthetic"),
        findings=(
            AuditFinding(
                control_id="TEST-NA-EXPORT-001",
                title="Contrôle non applicable",
                status=AuditStatus.PASS,
                applicability=Applicability.NOT_APPLICABLE,
                message="Hors périmètre explicite.",
            ),
        ),
    )

    payload = report.model_dump(mode="json")
    assert payload["findings"][0]["status"] == "NOT_APPLICABLE"
    assert payload["findings"][0]["applicability"] == "not_applicable"

    with ZipFile(BytesIO(render_docx(report))) as package:
        document = package.read("word/document.xml").decode("utf-8")
    assert "Contrôle non applicable" not in document
    assert "NON APPLICABLE" not in document
    assert "NOT_APPLICABLE" not in document
    assert "not_applicable" not in document

    workbook = load_workbook(BytesIO(render_xlsx(report)), read_only=True, data_only=True)
    enriched = list(workbook["Contrôles enrichis"].iter_rows(values_only=True))
    assert enriched[1][5:7] == ("not_applicable", "NOT_APPLICABLE")
