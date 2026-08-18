import re
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from uuid import UUID
from zipfile import ZipFile

import httpx
import pytest
from docx import Document
from openpyxl import load_workbook

from vysion.adapters.fortiguard import FortiGuardResult, FortiGuardStatus
from vysion.api.app import create_app
from vysion.audit.models import ExternalObservationStatus, PsirtObservation
from vysion.config import Settings

SYNTHETIC_CONFIG = b"""\
# Synthetic fixture created for Vysion v2
config system global
    set hostname "api-lab.example"
end
config system interface
    edit "wan1"
        set ip 192.0.2.20 255.255.255.0
        set role wan
        set allowaccess ping
    next
end
config system admin
    edit "secops"
        set two-factor fortitoken
    next
end
config user local
end
"""
REALISTIC_FIXTURE = Path(__file__).parents[1] / "fixtures" / "anonymized_fortigate_export.conf"
M3_IDS = (
    "FW-IMPLICIT-DENY-LOG-001",
    "FW-INTERNET-ALL-SERVICE-001",
    "FW-UTM-PROFILE-BINDING-001",
    "FW-VIP-EXTINTF-ANY-001",
    "FW-VSERVER-EXTINTF-ANY-001",
    "FW-SENSITIVE-PROTOCOL-DENY-001",
)
CONTROL_IDS = (
    "SYS-HOSTNAME-001",
    "NET-WAN-MGMT-001",
    "IAM-ADMIN-MFA-001",
    "IAM-LOCAL-USER-MFA-001",
    "IAM-DEFAULT-ADMIN-001",
    "IAM-GUEST-ACCOUNT-001",
    *M3_IDS,
    "VPN-SSL-001",
    "VPN-IKEV2-001",
    "VPN-DH-001",
    "VPN-CRYPTO-001",
    "UTM-LICENSE-001",
    "UTM-AUTOUPDATE-001",
    "UTM-DNSFILTER-001",
    "UTM-WEBFILTER-001",
    "UTM-ANTIVIRUS-001",
    "UTM-IPS-001",
    "UTM-APPCONTROL-001",
    "IAM-LDAPS-001",
    "EXT-PSIRT-001",
    "SYS-BACKUP-AUTO-001",
    "CFG-REF-INTEGRITY-001",
    "SYS-AUTO-INSTALL-USB-001",
    "SYS-FORTIMANAGER-SYNC-001",
    "SYS-FORTIANALYZER-SYNC-001",
    "SYS-ADMIN-HTTPS-PORT-001",
    "NET-SIP-ALG-001",
    "HA-SESSION-PICKUP-001",
    "HA-HEARTBEAT-REDUNDANCY-001",
    "HA-OVERRIDE-001",
    "HA-CABLING-REDUNDANCY-001",
    "UTM-FORTISANDBOX-CLOUD-001",
    "UTM-FORTIGUARD-ANYCAST-001",
    "NET-SDWAN-USAGE-001",
    "FW-BY-SEQUENCE-USAGE-001",
    "UTM-MAIL-FILTER-USAGE-001",
    "FW-SSL-SSH-PROFILE-001",
    "CFG-UNUSED-SERVICE-001",
    "IAM-LEGACY-ADMIN-001",
    "IAM-LEGACY-PKI-REMOVAL-001",
    "IAM-LEGACY-PKI-PRESENCE-001",
    "NET-LEGACY-ADMIN-LOOPBACK-001",
    "DNS-LEGACY-DATABASE-001",
    "NET-GEO-IP-USAGE-001",
    "NET-RFC6890-BLACKHOLE-001",
    "FW-LEGACY-SCHEDULE-INVENTORY-001",
    "WIFI-FORTIAP-OBSOLETE-001",
    "WIFI-SSID-LIMIT-001",
    "WIFI-RADIO2-40MHZ-001",
    "WIFI-DARRP-001",
    "WIFI-FREQUENCY-HANDOFF-001",
    "WIFI-TIM-001",
    "WIFI-BAND-001",
    "WIFI-CHANNELS-001",
    "WIFI-SHORT-GUARD-INTERVAL-001",
)

M3_FAIL_CONFIG = b"""\
config system interface
    edit "port1"
        set role lan
        set allowaccess ping
    next
    edit "wan1"
        set role wan
        set allowaccess ping
    next
end
config log setting
    set fwpolicy-implicit-log enable
end
config firewall service custom
    edit "sensitive"
        set tcp-portrange 88 389 636 445 137-139
        set udp-portrange 88 389 1812-1813 137-139
    next
end
config firewall policy
    edit 1
        set status enable
        set srcintf "port1"
        set dstintf "wan1"
        set srcaddr "all"
        set dstaddr "all"
        set action accept
        set schedule "always"
        set service "ALL"
        set logtraffic all
    next
    edit 2
        set status enable
        set srcintf "port1"
        set dstintf "wan1"
        set srcaddr "all"
        set dstaddr "all"
        set action accept
        set schedule "always"
        set service "HTTPS"
        set logtraffic utm
        set utm-status enable
    next
    edit 3
        set status enable
        set srcintf "port1"
        set dstintf "wan1"
        set srcaddr "all"
        set dstaddr "all"
        set action accept
        set schedule "always"
        set service "sensitive"
        set logtraffic all
    next
end
config firewall vip
    edit "bad-vip"
        set extintf "any"
        set extip 198.51.100.20
        set mappedip "10.0.0.20"
    next
    edit "bad-vserver"
        set type server-load-balance
        set extintf "any"
        config realservers
            edit 1
                set ip 10.0.0.21
                set port 443
            next
        end
    next
end
"""


class AvailableFortiGuard:
    async def check(self) -> FortiGuardResult:
        return FortiGuardResult(status=FortiGuardStatus.AVAILABLE, detail="fixture")

    async def check_psirt(self, version: str) -> PsirtObservation:
        return PsirtObservation(
            status=ExternalObservationStatus.UNKNOWN,
            fortios_version=version,
            source="https://www.fortiguard.com/psirt",
            ruleset_id="fortiguard-psirt-critical-high",
            ruleset_version="2026-08-13",
            observed_at=datetime.now(UTC),
            complete=False,
        )


class ForbiddenPsirtFortiGuard(AvailableFortiGuard):
    async def check_psirt(self, version: str) -> PsirtObservation:
        raise AssertionError(f"unexpected PSIRT collection for {version}")


class CorrelatedPsirtFortiGuard(AvailableFortiGuard):
    def __init__(self) -> None:
        self.versions: list[str] = []

    async def check_psirt(self, version: str) -> PsirtObservation:
        self.versions.append(version)
        return PsirtObservation(
            status=ExternalObservationStatus.PASS,
            fortios_version=version,
            source="https://www.fortiguard.com/psirt",
            ruleset_id="fortiguard-psirt-critical-high",
            ruleset_version="2026-08-13",
            observed_at=datetime.now(UTC),
            complete=True,
        )


class MalformedPsirtFortiGuard(AvailableFortiGuard):
    async def check_psirt(self, version: str) -> PsirtObservation:
        return {"status": "PASS"}  # type: ignore[return-value]


def api_client(app) -> httpx.AsyncClient:
    return httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app),
        base_url="http://vysion.test",
    )


@pytest.mark.asyncio
async def test_api_does_not_collect_psirt_without_firmware_version(
    tmp_path: Path,
) -> None:
    app = create_app(
        settings=Settings(report_directory=tmp_path),
        fortiguard=ForbiddenPsirtFortiGuard(),
    )
    raw = "config system global\n    set hostname edge\nend\n"

    async with api_client(app) as client:
        response = await client.post(
            "/api/audits",
            files={"configuration": ("synthetic.conf", raw, "text/plain")},
        )

    assert response.status_code == 201
    payload = response.json()
    assert payload["context"]["psirt"] is None
    psirt = next(item for item in payload["findings"] if item["control_id"] == "EXT-PSIRT-001")
    assert psirt["status"] == "UNKNOWN"


@pytest.mark.asyncio
async def test_api_collects_psirt_and_registers_correlated_finding(tmp_path: Path) -> None:
    fortiguard = CorrelatedPsirtFortiGuard()
    app = create_app(settings=Settings(report_directory=tmp_path), fortiguard=fortiguard)
    raw = (
        "#config-version=FGT60E-7.2.9-FW-build1-1:opmode=0\n"
        "config system global\n    set hostname edge\nend\n"
    )

    async with api_client(app) as client:
        response = await client.post(
            "/api/audits",
            files={"configuration": ("synthetic.conf", raw, "text/plain")},
        )

    assert response.status_code == 201
    payload = response.json()
    assert fortiguard.versions == ["7.2.9"]
    assert payload["context"]["psirt"]["fortios_version"] == "7.2.9"
    psirt = next(item for item in payload["findings"] if item["control_id"] == "EXT-PSIRT-001")
    assert psirt["status"] == "PASS"


@pytest.mark.asyncio
async def test_api_rejects_malformed_psirt_observation_without_breaking_registry(
    tmp_path: Path,
) -> None:
    app = create_app(
        settings=Settings(report_directory=tmp_path),
        fortiguard=MalformedPsirtFortiGuard(),
    )
    raw = (
        "#config-version=FGT60E-7.2.9-FW-build1-1:opmode=0\n"
        "config system global\n    set hostname edge\nend\n"
    )

    async with api_client(app) as client:
        response = await client.post(
            "/api/audits",
            files={"configuration": ("synthetic.conf", raw, "text/plain")},
        )

    assert response.status_code == 201
    payload = response.json()
    assert payload["context"]["psirt"]["status"] == "ERROR"
    findings = {item["control_id"]: item for item in payload["findings"]}
    psirt = findings["EXT-PSIRT-001"]
    assert psirt["status"] == "UNKNOWN"
    assert findings["SYS-BACKUP-AUTO-001"]["status"] == "UNKNOWN"
    assert findings["CFG-REF-INTEGRITY-001"]["status"] == "UNKNOWN"
    assert findings["SYS-AUTO-INSTALL-USB-001"]["status"] == "UNKNOWN"
    assert findings["SYS-FORTIMANAGER-SYNC-001"]["status"] == "UNKNOWN"
    assert findings["SYS-FORTIANALYZER-SYNC-001"]["status"] == "UNKNOWN"


@pytest.mark.asyncio
async def test_api_accepts_anonymized_realistic_fortigate_export(tmp_path: Path) -> None:
    app = create_app(
        settings=Settings(report_directory=tmp_path),
        fortiguard=CorrelatedPsirtFortiGuard(),
    )

    async with api_client(app) as client:
        response = await client.post(
            "/api/audits",
            data={
                "utm_license": "true",
                "context_source": "integration-fixture",
                "context_method": "explicit-test-context",
            },
            files={
                "configuration": (
                    "anonymized-fortigate.conf",
                    REALISTIC_FIXTURE.read_bytes(),
                    "text/plain",
                )
            },
        )

    assert response.status_code == 201
    findings = response.json()["findings"]
    expected_statuses = ["PASS"] * len(CONTROL_IDS)
    expected_statuses[CONTROL_IDS.index("VPN-SSL-001")] = "NOT_APPLICABLE"
    expected_statuses[CONTROL_IDS.index("FW-SSL-SSH-PROFILE-001")] = "NOT_APPLICABLE"
    expected_statuses[CONTROL_IDS.index("SYS-BACKUP-AUTO-001")] = "UNKNOWN"
    expected_statuses[CONTROL_IDS.index("CFG-REF-INTEGRITY-001")] = "UNKNOWN"
    for control_id in (
        "SYS-AUTO-INSTALL-USB-001",
        "SYS-FORTIMANAGER-SYNC-001",
        "SYS-FORTIANALYZER-SYNC-001",
        "SYS-ADMIN-HTTPS-PORT-001",
        "NET-SIP-ALG-001",
        "HA-SESSION-PICKUP-001",
        "HA-HEARTBEAT-REDUNDANCY-001",
        "HA-OVERRIDE-001",
        "HA-CABLING-REDUNDANCY-001",
        "UTM-FORTISANDBOX-CLOUD-001",
        "UTM-FORTIGUARD-ANYCAST-001",
        "NET-SDWAN-USAGE-001",
        "IAM-LEGACY-ADMIN-001",
        "IAM-LEGACY-PKI-REMOVAL-001",
        "IAM-LEGACY-PKI-PRESENCE-001",
        "NET-LEGACY-ADMIN-LOOPBACK-001",
        "DNS-LEGACY-DATABASE-001",
        "NET-GEO-IP-USAGE-001",
        "NET-RFC6890-BLACKHOLE-001",
        "FW-LEGACY-SCHEDULE-INVENTORY-001",
        *CONTROL_IDS[-9:],
    ):
        expected_statuses[CONTROL_IDS.index(control_id)] = "UNKNOWN"
    expected_statuses[CONTROL_IDS.index("CFG-UNUSED-SERVICE-001")] = "UNKNOWN"
    expected_statuses[CONTROL_IDS.index("FW-BY-SEQUENCE-USAGE-001")] = "FAIL"
    assert [finding["status"] for finding in findings] == expected_statuses
    assert all(
        finding["evidence_items"]
        and all(item["certainty"] == "certain" for item in finding["evidence_items"])
        for finding in findings
        if finding["status"] != "UNKNOWN"
    )
    assert tuple(finding["control_id"] for finding in findings) == CONTROL_IDS


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "configuration",
    [
        b"""config system interface extra
    edit "wan-hidden"
        set allowaccess ssh
    next
end
config system interface
    edit "wan1"
        set allowaccess ping https
    next
end
""",
        b"""config system interface
    set allowaccess ssh
    edit "wan1"
        set allowaccess ping https
    next
end
""",
    ],
)
async def test_api_rejects_ambiguous_audited_evidence(
    tmp_path: Path,
    configuration: bytes,
) -> None:
    app = create_app(
        settings=Settings(report_directory=tmp_path),
        fortiguard=AvailableFortiGuard(),
    )

    async with api_client(app) as client:
        response = await client.post(
            "/api/audits",
            files={"configuration": ("ambiguous.conf", configuration, "text/plain")},
        )

    assert response.status_code == 422
    assert not list(tmp_path.glob("*.json"))


@pytest.mark.asyncio
async def test_api_never_reports_pass_from_a_stale_hostname_proof(tmp_path: Path) -> None:
    app = create_app(
        settings=Settings(report_directory=tmp_path),
        fortiguard=AvailableFortiGuard(),
    )
    configuration = b"""config system global
    set hostname safe.example
    set hostname
end
"""

    async with api_client(app) as client:
        response = await client.post(
            "/api/audits",
            files={"configuration": ("stale-proof.conf", configuration, "text/plain")},
        )

    assert response.status_code == 201
    assert [finding["status"] for finding in response.json()["findings"]] == [
        "UNKNOWN",
    ] * len(CONTROL_IDS)


@pytest.mark.asyncio
async def test_api_stores_a_typed_json_report_under_uuid_and_serves_it(
    tmp_path: Path,
) -> None:
    now = datetime(2026, 8, 12, 12, 0, tzinfo=UTC)
    settings = Settings(report_directory=tmp_path, report_ttl_seconds=60)
    app = create_app(settings=settings, fortiguard=AvailableFortiGuard(), clock=lambda: now)

    async with api_client(app) as client:
        response = await client.post(
            "/api/audits",
            files={"configuration": ("synthetic.conf", SYNTHETIC_CONFIG, "text/plain")},
        )

        assert response.status_code == 201
        assert response.headers["cache-control"] == "no-store, private"
        payload = response.json()
        report_id = UUID(payload["report_id"], version=4)
        assert payload["expires_at"] == (now + timedelta(seconds=60)).isoformat().replace(
            "+00:00", "Z"
        )
        assert payload["fortiguard"]["status"] == "AVAILABLE"
        statuses = {finding["control_id"]: finding["status"] for finding in payload["findings"]}
        assert {
            control_id
            for control_id, finding_status in statuses.items()
            if finding_status == "PASS"
        } == {
            "SYS-HOSTNAME-001",
            "NET-WAN-MGMT-001",
            "IAM-ADMIN-MFA-001",
            "IAM-DEFAULT-ADMIN-001",
            "IAM-GUEST-ACCOUNT-001",
        }
        assert statuses["IAM-LOCAL-USER-MFA-001"] == "NOT_APPLICABLE"
        assert payload["findings"][3]["applicability"] == "not_applicable"
        assert all(
            finding["applicability"] == "unknown"
            for finding in payload["findings"]
            if finding["status"] == "UNKNOWN"
        )
        assert all(
            finding["priority"]
            == (
                "P1"
                if finding["control_id"]
                in {
                    "SYS-BACKUP-AUTO-001",
                    "CFG-REF-INTEGRITY-001",
                    "SYS-AUTO-INSTALL-USB-001",
                    "SYS-FORTIMANAGER-SYNC-001",
                    "SYS-FORTIANALYZER-SYNC-001",
                    "SYS-ADMIN-HTTPS-PORT-001",
                    "NET-SIP-ALG-001",
                    "HA-SESSION-PICKUP-001",
                    "HA-HEARTBEAT-REDUNDANCY-001",
                    "HA-OVERRIDE-001",
                    "HA-CABLING-REDUNDANCY-001",
                    "UTM-FORTISANDBOX-CLOUD-001",
                    "UTM-FORTIGUARD-ANYCAST-001",
                    "NET-SDWAN-USAGE-001",
                    "FW-BY-SEQUENCE-USAGE-001",
                    "UTM-MAIL-FILTER-USAGE-001",
                    "FW-SSL-SSH-PROFILE-001",
                    "CFG-UNUSED-SERVICE-001",
                    "IAM-LEGACY-ADMIN-001",
                    "IAM-LEGACY-PKI-REMOVAL-001",
                    "IAM-LEGACY-PKI-PRESENCE-001",
                    "NET-LEGACY-ADMIN-LOOPBACK-001",
                    "DNS-LEGACY-DATABASE-001",
                    "NET-GEO-IP-USAGE-001",
                    "NET-RFC6890-BLACKHOLE-001",
                    "FW-LEGACY-SCHEDULE-INVENTORY-001",
                    "WIFI-FORTIAP-OBSOLETE-001",
                    "WIFI-SSID-LIMIT-001",
                    "WIFI-RADIO2-40MHZ-001",
                    "WIFI-DARRP-001",
                    "WIFI-FREQUENCY-HANDOFF-001",
                    "WIFI-TIM-001",
                    "WIFI-BAND-001",
                    "WIFI-CHANNELS-001",
                    "WIFI-SHORT-GUARD-INTERVAL-001",
                }
                else "P0"
            )
            for finding in payload["findings"][1:]
        )
        assert (tmp_path / f"{report_id}.json").is_file()

        stored = await client.get(f"/api/reports/{report_id}.json")
        assert stored.status_code == 200
        assert stored.headers["content-type"] == "application/json"
        assert stored.headers["cache-control"] == "no-store, private"
        assert stored.json() == payload


@pytest.mark.asyncio
async def test_api_exposes_m3_failures_without_hiding_certain_violations(
    tmp_path: Path,
) -> None:
    app = create_app(
        settings=Settings(report_directory=tmp_path),
        fortiguard=AvailableFortiGuard(),
    )

    async with api_client(app) as client:
        response = await client.post(
            "/api/audits",
            files={"configuration": ("m3-fail.conf", M3_FAIL_CONFIG, "text/plain")},
        )

    assert response.status_code == 201
    findings = {finding["control_id"]: finding for finding in response.json()["findings"]}
    assert findings["FW-IMPLICIT-DENY-LOG-001"]["status"] == "PASS"
    for control_id in M3_IDS[1:]:
        finding = findings[control_id]
        assert finding["status"] == "FAIL"
        assert finding["applicability"] == "applicable"
        assert finding["evidence_items"]


@pytest.mark.asyncio
async def test_api_json_docx_xlsx_preserve_all_control_ids(
    tmp_path: Path,
) -> None:
    app = create_app(
        settings=Settings(report_directory=tmp_path, report_ttl_seconds=60),
        fortiguard=AvailableFortiGuard(),
    )

    async with api_client(app) as client:
        created = await client.post(
            "/api/audits",
            files={
                "configuration": (
                    "anonymized-fortigate.conf",
                    REALISTIC_FIXTURE.read_bytes(),
                    "text/plain",
                )
            },
        )
        assert created.status_code == 201
        report_id = created.json()["report_id"]
        json_ids = [finding["control_id"] for finding in created.json()["findings"]]

        docx_response = await client.get(f"/api/reports/{report_id}.docx")
        xlsx_response = await client.get(f"/api/reports/{report_id}.xlsx")

    expected_ids = json_ids
    assert tuple(expected_ids) == CONTROL_IDS
    assert len(set(expected_ids)) == len(CONTROL_IDS)

    document = Document(BytesIO(docx_response.content))
    docx_text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )
    docx_ids = re.findall(r"\b[A-Z][A-Z0-9]+(?:-[A-Z0-9]+)+-\d{3}\b", docx_text)
    assert set(docx_ids) == set(expected_ids)

    workbook = load_workbook(BytesIO(xlsx_response.content), read_only=True, data_only=True)
    xlsx_ids = [
        row[0] for row in workbook["Contrôles enrichis"].iter_rows(min_row=2, values_only=True)
    ]
    assert xlsx_ids == expected_ids


@pytest.mark.asyncio
async def test_api_round_trips_explicit_operator_context_without_false_defaults(
    tmp_path: Path,
) -> None:
    app = create_app(
        settings=Settings(report_directory=tmp_path),
        fortiguard=AvailableFortiGuard(),
    )

    async with api_client(app) as client:
        response = await client.post(
            "/api/audits",
            data={
                "selected_wans": '["wan1"]',
                "client": "Client synthétique",
                "site": "Paris-lab",
                "ha": "true",
                "ha_cabling_redundancy": "true",
                "utm_license": "false",
                "context_source": "operator-form",
                "context_operator": "analyst@example.invalid",
                "context_method": "manual-selection",
            },
            files={"configuration": ("synthetic.conf", SYNTHETIC_CONFIG, "text/plain")},
        )

        assert response.status_code == 201
        payload = response.json()
        assert payload["context"] == {
            "selected_wans": ["wan1"],
            "operator_provenance": {
                "source": "operator-form",
                "operator": "analyst@example.invalid",
                "captured_at": None,
                "method": "manual-selection",
            },
            "client": "Client synthétique",
            "site": "Paris-lab",
            "operator_comment": None,
            "ha": True,
            "ha_cabling_redundancy": True,
            "mpls": None,
            "utm_license": False,
            "psirt": None,
        }
        stored = await client.get(f"/api/reports/{payload['report_id']}.json")
        assert stored.json() == payload


@pytest.mark.asyncio
async def test_api_keeps_contradictory_selected_wan_unknown_and_traceable(tmp_path: Path) -> None:
    app = create_app(
        settings=Settings(report_directory=tmp_path),
        fortiguard=AvailableFortiGuard(),
    )

    async with api_client(app) as client:
        response = await client.post(
            "/api/audits",
            data={"selected_wans": '["missing-wan"]'},
            files={"configuration": ("synthetic.conf", SYNTHETIC_CONFIG, "text/plain")},
        )

    assert response.status_code == 201
    finding = next(
        item for item in response.json()["findings"] if item["control_id"] == "NET-WAN-MGMT-001"
    )
    assert finding["status"] == "UNKNOWN"
    assert finding["applicability"] == "unknown"
    assert "missing-wan" in " ".join(finding["evidence"])


@pytest.mark.asyncio
async def test_api_generates_docx_from_the_stored_typed_report(tmp_path: Path) -> None:
    now = datetime(2026, 8, 12, 12, 0, tzinfo=UTC)
    app = create_app(
        settings=Settings(report_directory=tmp_path, report_ttl_seconds=60),
        fortiguard=AvailableFortiGuard(),
        clock=lambda: now,
    )

    async with api_client(app) as client:
        created = await client.post(
            "/api/audits",
            files={"configuration": ("synthetic.conf", SYNTHETIC_CONFIG, "text/plain")},
        )
        report_id = UUID(created.json()["report_id"])

        response = await client.get(f"/api/reports/{report_id}.docx")

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert response.headers["cache-control"] == "no-store, private"
    assert response.headers["content-disposition"] == (
        f'attachment; filename="vysion-{report_id}.docx"'
    )
    with ZipFile(BytesIO(response.content)) as package:
        document = package.read("word/document.xml").decode("utf-8")
    assert "Rapport d’audit Vysion" in document
    assert "synthetic.conf" in document
    assert "SYS-HOSTNAME-001" in document
    assert "IAM-LOCAL-USER-MFA-001" in document
    assert "IAM-DEFAULT-ADMIN-001" in document
    assert "IAM-GUEST-ACCOUNT-001" in document
    assert "AVAILABLE" in document


@pytest.mark.asyncio
async def test_api_generates_xlsx_from_the_stored_typed_report(tmp_path: Path) -> None:
    now = datetime(2026, 8, 12, 12, 0, tzinfo=UTC)
    app = create_app(
        settings=Settings(report_directory=tmp_path, report_ttl_seconds=60),
        fortiguard=AvailableFortiGuard(),
        clock=lambda: now,
    )

    async with api_client(app) as client:
        created = await client.post(
            "/api/audits",
            files={"configuration": ("synthetic.conf", SYNTHETIC_CONFIG, "text/plain")},
        )
        report_id = UUID(created.json()["report_id"])

        response = await client.get(f"/api/reports/{report_id}.xlsx")

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["cache-control"] == "no-store, private"
    assert response.headers["content-disposition"] == (
        f'attachment; filename="vysion-{report_id}.xlsx"'
    )
    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
    assert workbook.sheetnames[:3] == ["Synthèse", "Contrôles", "Contrôles enrichis"]
    assert {
        "Audit configuration",
        "Actions sans accord",
        "Actions avec accord",
        "Statistiques",
        "Comptes",
        "Métadonnées équipement",
    } <= set(workbook.sheetnames)
    summary = {
        str(key): value
        for key, value in workbook["Synthèse"].iter_rows(
            min_row=1,
            max_col=2,
            values_only=True,
        )
    }
    assert summary["Source"] == "synthetic.conf"
    assert summary["FortiGuard"] == "AVAILABLE"
    controls = list(workbook["Contrôles"].iter_rows(values_only=True))
    assert controls[0] == ("Contrôle", "Titre", "Statut", "Constat", "Risque", "Recommandation")
    assert controls[1][0] == "SYS-HOSTNAME-001"
    assert controls[1][2] == "PASS"
    assert {row[0] for row in controls[1:]} >= {
        "IAM-LOCAL-USER-MFA-001",
        "IAM-DEFAULT-ADMIN-001",
        "IAM-GUEST-ACCOUNT-001",
        "FW-IMPLICIT-DENY-LOG-001",
        "FW-INTERNET-ALL-SERVICE-001",
        "FW-UTM-PROFILE-BINDING-001",
        "FW-VIP-EXTINTF-ANY-001",
        "FW-VSERVER-EXTINTF-ANY-001",
        "FW-SENSITIVE-PROTOCOL-DENY-001",
    }
    enriched = list(workbook["Contrôles enrichis"].iter_rows(values_only=True))
    assert enriched[0][2:7] == (
        "Catégorie",
        "Priorité",
        "Sévérité",
        "Applicabilité",
        "Statut",
    )
    assert enriched[1][0] == "SYS-HOSTNAME-001"
    assert enriched[1][6] == "PASS"


@pytest.mark.asyncio
async def test_expired_report_is_deleted_and_returns_404(tmp_path: Path) -> None:
    current = datetime(2026, 8, 12, 12, 0, tzinfo=UTC)
    settings = Settings(report_directory=tmp_path, report_ttl_seconds=60)
    app = create_app(settings=settings, fortiguard=AvailableFortiGuard(), clock=lambda: current)

    async with api_client(app) as client:
        created = (
            await client.post(
                "/api/audits",
                files={"configuration": ("synthetic.conf", SYNTHETIC_CONFIG, "text/plain")},
            )
        ).json()
        report_path = tmp_path / f"{created['report_id']}.json"

        current += timedelta(seconds=61)
        expired = await client.get(f"/api/reports/{created['report_id']}.json")

        assert expired.status_code == 404
        assert not report_path.exists()


@pytest.mark.asyncio
async def test_health_reports_application_readiness(tmp_path: Path) -> None:
    app = create_app(
        settings=Settings(report_directory=tmp_path),
        fortiguard=AvailableFortiGuard(),
    )

    async with api_client(app) as client:
        response = await client.get("/api/health")

    assert response.status_code == 200
    assert response.json() == {"status": "ok", "service": "vysion", "version": "2.2.0-dev"}
