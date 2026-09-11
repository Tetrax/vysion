from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from vysion.audit.models import (
    Applicability,
    AuditFinding,
    AuditStatus,
    DeviceIdentity,
    RiskAssessment,
)
from vysion.reports.business_text import (
    business_risk_for,
    business_text_for,
    client_result_for,
)
from vysion.reports.json_report import JsonAuditReport
from vysion.reports.presentation import AuditPresentation, build_presentation
from vysion.reports.v1_document import V1_DOCUMENT_SECTIONS, V1DocumentSlot

_TEMPLATE = Path(__file__).with_name("templates") / "generique.docx"

_DOMAIN_ORDER = (
    "system",
    "administration",
    "network",
    "cluster",
    "vpn",
    "utm",
    "wifi",
)
_DOMAIN_TITLES = {
    "system": "Audit - Système",
    "administration": "Audit - Administration et comptes",
    "network": "Audit - Réseaux et flux",
    "cluster": "Audit - Cluster",
    "vpn": "Audit - Connexions distantes : VPN",
    "utm": "Audit - Profils de sécurité UTM",
    "wifi": "Audit - Wi-Fi et FortiAP",
}

# The V1 report uses these four business levels.  V2 findings may carry either
# the historical French labels or the generic English labels from the model.
_IMPACT_ALIASES = {
    "critical": "CRITIQUE",
    "critique": "CRITIQUE",
    "high": "GRAVE",
    "grave": "GRAVE",
    "medium": "SIGNIFICATIF",
    "significatif": "SIGNIFICATIF",
    "low": "NEGLIGEABLE",
    "negligeable": "NEGLIGEABLE",
}
_LIKELIHOOD_ALIASES = {
    "almost certain": "QUASI CERTAIN",
    "quasi certain": "QUASI CERTAIN",
    "very likely": "TRÈS VRAISEMBLABLE",
    "très vraisemblable": "TRÈS VRAISEMBLABLE",
    "likely": "VRAISEMBLABLE",
    "vraisemblable": "VRAISEMBLABLE",
    "unlikely": "PEU VRAISEMBLABLE",
    "peu vraisemblable": "PEU VRAISEMBLABLE",
}
_CORRECTION_ALIASES = {
    "complex": "COMPLEXE",
    "complexe": "COMPLEXE",
    "reasonable": "RAISONNABLE",
    "raisonnable": "RAISONNABLE",
    "simple": "SIMPLE",
}
_RISK_MATRIX = {
    "CRITIQUE": {
        "PEU VRAISEMBLABLE": "MOYEN",
        "VRAISEMBLABLE": "ÉLEVÉ",
        "TRÈS VRAISEMBLABLE": "TRÈS ÉLEVÉ",
        "QUASI CERTAIN": "TRÈS ÉLEVÉ",
    },
    "GRAVE": {
        "PEU VRAISEMBLABLE": "FAIBLE",
        "VRAISEMBLABLE": "ÉLEVÉ",
        "TRÈS VRAISEMBLABLE": "ÉLEVÉ",
        "QUASI CERTAIN": "TRÈS ÉLEVÉ",
    },
    "SIGNIFICATIF": {
        "PEU VRAISEMBLABLE": "FAIBLE",
        "VRAISEMBLABLE": "MOYEN",
        "TRÈS VRAISEMBLABLE": "ÉLEVÉ",
        "QUASI CERTAIN": "ÉLEVÉ",
    },
    "NEGLIGEABLE": {
        "PEU VRAISEMBLABLE": "FAIBLE",
        "VRAISEMBLABLE": "FAIBLE",
        "TRÈS VRAISEMBLABLE": "MOYEN",
        "QUASI CERTAIN": "MOYEN",
    },
}
_RISK_COLORS = {
    "TRÈS ÉLEVÉ": "C00000",
    "ÉLEVÉ": "F4B183",
    "MOYEN": "FFFF00",
    "FAIBLE": "C6E0B4",
}
_IMPACT_COLORS = {
    "CRITIQUE": "C00000",
    "GRAVE": "F4B183",
    "SIGNIFICATIF": "FFFF00",
    "NEGLIGEABLE": "C6E0B4",
}
_LIKELIHOOD_COLORS = {
    "QUASI CERTAIN": "C00000",
    "TRÈS VRAISEMBLABLE": "F4B183",
    "VRAISEMBLABLE": "FFFF00",
    "PEU VRAISEMBLABLE": "C6E0B4",
}
_CORRECTION_COLORS = {
    "COMPLEXE": "C00000",
    "RAISONNABLE": "F4B183",
    "SIMPLE": "C6E0B4",
}


def _display(value: object, default: str = "Non renseigné") -> str:
    if value is None or value == "":
        return default
    return str(value)


def _clean_client_text(value: str | None) -> str:
    """Keep business prose while removing internal engine vocabulary."""

    text = _display(value)
    text = re.sub(r"\b[A-Z][A-Z0-9]+(?:-[A-Z0-9]+)+-\d{3}\b", "ce contrôle", text)
    text = re.sub(r"\b(?:policy|règle)\s+\d+\b", "la règle concernée", text, flags=re.IGNORECASE)
    text = re.sub(r"\bligne\s+\d+\b", "", text, flags=re.IGNORECASE)
    text = re.sub(
        r"\b(?:rule_version|ruleset-version|ruleset|provenance|parser|namespaces?|evidence)\b(?:\s*[:=-]?\s*[\w.-]+)?",
        "",
        text,
        flags=re.IGNORECASE,
    )
    text = re.sub(r"\bcertain\b", "confirmé", text, flags=re.IGNORECASE)
    text = re.sub(r"\bcertaine\b", "confirmée", text, flags=re.IGNORECASE)
    text = re.sub(r"\bcertains\b", "confirmés", text, flags=re.IGNORECASE)
    text = re.sub(r"\bcertaines\b", "confirmées", text, flags=re.IGNORECASE)
    text = re.sub(r"\buncertain\b", "non confirmé", text, flags=re.IGNORECASE)
    text = re.sub(r"\bambiguous\b", "incomplet", text, flags=re.IGNORECASE)
    text = re.sub(r"\bdefaulted\b", "par défaut", text, flags=re.IGNORECASE)
    text = re.sub(r"\bproven\b", "confirmé", text, flags=re.IGNORECASE)
    text = re.sub(r"\s{2,}", " ", text).strip(" .;:-")
    text = re.sub(r"^[=+@]", "", text)
    return text or "ce contrôle"


def _replace_everywhere(document: Any, replacements: dict[str, str]) -> None:
    containers = [document]
    for section in document.sections:
        containers.extend((section.header, section.footer))

    for container in containers:
        paragraphs = list(container.paragraphs)
        for table in container.tables:
            paragraphs.extend(cell.paragraphs for row in table.rows for cell in row.cells)
        for paragraph_group in paragraphs:
            if not isinstance(paragraph_group, list):
                paragraph_group = [paragraph_group]
            for paragraph in paragraph_group:
                for run in paragraph.runs:
                    for old, new in replacements.items():
                        if old in run.text:
                            run.text = run.text.replace(old, new)
                replacement_text = paragraph.text
                for old, new in replacements.items():
                    replacement_text = replacement_text.replace(old, new)
                if replacement_text != paragraph.text:
                    # Template prose can split a client-specific phrase over
                    # several runs.  This fallback is limited to those static
                    # template paragraphs; generated report formatting is
                    # applied separately below.
                    paragraph.text = replacement_text


def _shade(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_cell_text(cell: Any, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _repeat_table_header(row: Any) -> None:
    """Repeat a table's header row when the table crosses a page break."""

    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _color_status(run, status: AuditStatus) -> None:
    run.bold = True
    run.font.color.rgb = (
        RGBColor(
            0x00,
            0xB0,
            0x50,
        )
        if status is AuditStatus.PASS
        else RGBColor(
            0xCC,
            0x00,
            0x00,
        )
        if status is AuditStatus.FAIL
        else RGBColor(0xFF, 0xC0, 0x00)
    )


def _domain_for(finding: AuditFinding) -> str:
    control_id = finding.control_id.upper()
    if control_id.startswith("HA-"):
        return "cluster"
    if control_id.startswith("VPN-"):
        return "vpn"
    if control_id.startswith(("UTM-",)):
        return "utm"
    if control_id.startswith(("WIFI-", "FAP-")):
        return "wifi"
    if control_id.startswith(("IAM-", "NET-WAN-MGMT-")):
        return "administration"
    if control_id.startswith(("NET-LEGACY-ADMIN-", "DNS-LEGACY-")):
        return "administration"
    if control_id.startswith(("NET-", "FW-")):
        return "network"
    return "system"


def _point_audited(finding: AuditFinding) -> str:
    business_text = business_text_for(finding)
    if business_text is not None:
        return business_text.point
    return (
        "Évaluation de ce point de configuration à partir des éléments réellement "
        f"observés : {_clean_client_text(finding.display_name or finding.title)}."
    )


def _result_explanation(finding: AuditFinding) -> str:
    if finding.control_id in {
        "SYS-FORTIMANAGER-SYNC-001",
        "SYS-FORTIANALYZER-SYNC-001",
    }:
        target = (
            "FortiManager" if finding.control_id == "SYS-FORTIMANAGER-SYNC-001" else "FortiAnalyzer"
        )
        if finding.status is AuditStatus.PASS:
            return (
                f"Le FortiGate est configuré avec un {target} explicite. "
                "L'état réel de synchronisation doit être vérifié."
            )
        return (
            f"La configuration {target} ne permet pas de conclure sur l'état réel "
            "de synchronisation."
        )
    if finding.status is AuditStatus.NOT_APPLICABLE:
        return "Ce point n'est pas applicable au périmètre analysé."
    return client_result_for(finding)


def _normalize_impact(value: str | None) -> str:
    return _IMPACT_ALIASES.get((value or "").casefold(), "SIGNIFICATIF")


def _normalize_likelihood(value: str | None) -> str:
    return _LIKELIHOOD_ALIASES.get((value or "").casefold(), "VRAISEMBLABLE")


def _normalize_correction(value: str | None) -> str:
    return _CORRECTION_ALIASES.get((value or "").casefold(), "SIMPLE")


def _risk_values(risk: RiskAssessment | None) -> tuple[str, str, str, str]:
    impact = _normalize_impact(risk.impact if risk else None)
    likelihood = _normalize_likelihood(risk.likelihood if risk else None)
    correction = _normalize_correction(risk.treatment if risk else None)
    level = _RISK_MATRIX[impact][likelihood]
    return likelihood, impact, level, correction


def _add_risk_table(
    document: Any,
    risk_number: int,
    finding: AuditFinding,
) -> dict[str, str]:
    business_text = business_text_for(finding)
    risk_assessment = business_risk_for(finding)
    likelihood, impact, risk_level, correction = _risk_values(risk_assessment)
    risk_point = (
        business_text.risk_point
        if business_text and business_text.risk_point
        else finding.display_name or finding.title
    )
    point = _clean_client_text(risk_point)
    description = _clean_client_text(
        business_text.risk_description
        if business_text and business_text.risk_description
        else risk_assessment.summary
        if risk_assessment
        else "La configuration présente une faiblesse de sécurité à corriger."
    )
    remediation = _clean_client_text(
        business_text.remediation
        if business_text and business_text.remediation
        else finding.remediation
        or finding.recommendation
        or "Mettre la configuration en conformité."
    )
    table = document.add_table(rows=7, cols=2)
    table.style = "Table Grid"
    labels = (
        (f"R{risk_number}", point),
        ("DESCRIPTION DU RISQUE", description),
        ("VRAISEMBLANCE", likelihood),
        ("IMPACT", impact),
        ("RISQUE", risk_level),
        ("CORRECTION", correction),
        ("REMEDIATION", remediation),
    )
    for index, (label, value) in enumerate(labels):
        _set_cell_text(table.cell(index, 0), label, bold=True)
        _set_cell_text(table.cell(index, 1), value, bold=index == 0)
        if label == "RISQUE":
            _shade(table.cell(index, 1), _RISK_COLORS[risk_level])
        elif label == "IMPACT":
            _shade(table.cell(index, 1), _IMPACT_COLORS[impact])
        elif label == "VRAISEMBLANCE":
            _shade(table.cell(index, 1), _LIKELIHOOD_COLORS[likelihood])
        elif label == "CORRECTION":
            _shade(table.cell(index, 1), _CORRECTION_COLORS[correction])
        if label in {"RISQUE", "IMPACT", "VRAISEMBLANCE", "CORRECTION"}:
            for run in table.cell(index, 1).paragraphs[0].runs:
                run.bold = label == "RISQUE"
                if (
                    (label == "RISQUE" and risk_level == "TRÈS ÉLEVÉ")
                    or (label == "IMPACT" and impact == "CRITIQUE")
                    or (label == "VRAISEMBLANCE" and likelihood == "QUASI CERTAIN")
                    or (label == "CORRECTION" and correction == "COMPLEXE")
                ):
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row in table.rows:
        row.cells[0].width = Cm(3.8)
        row.cells[1].width = Cm(13)
    document.add_paragraph("")
    return {
        "id": f"R{risk_number}",
        "point": point,
        "description": description,
        "likelihood": likelihood,
        "impact": impact,
        "risk": risk_level,
        "correction": correction,
    }


def _presentation_for(report: JsonAuditReport) -> AuditPresentation:
    return report.presentation or build_presentation(report.findings)


def _add_presentation_summary(document: Any, presentation: AuditPresentation) -> None:
    document.add_heading("Périmètre de comparaison V1/V2", level=2)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    _set_cell_text(table.cell(0, 0), "Indicateur", bold=True)
    _set_cell_text(table.cell(0, 1), "Valeur", bold=True)
    for label, value in (
        ("Points métier V1 comparables", presentation.business_control_count),
        ("Contrôles moteur exécutés", presentation.engine_control_count),
        ("Sous-vérifications détaillées", presentation.split_extra_finding_count),
        ("Contrôles complémentaires", presentation.v2_only_control_count),
    ):
        cells = table.add_row().cells
        _set_cell_text(cells[0], label)
        _set_cell_text(cells[1], str(value))
    for line in presentation.explanation_lines:
        document.add_paragraph(line)
    document.add_paragraph("")


def _add_characteristics(document: Any, report: JsonAuditReport) -> None:
    identity: DeviceIdentity | None = report.device_identity
    context = report.context
    ha_observed = any(
        finding.control_id.startswith("HA-")
        and any(
            item.section and item.section.casefold() == "system ha"
            for item in finding.evidence_items
        )
        for finding in report.findings
    )
    cluster = (
        "Oui"
        if ha_observed
        else _display(context.ha)
        if context.ha is not None
        else "Non renseigné"
    )
    if context.utm_license is True:
        license_state = "Valide"
    elif context.utm_license is False:
        license_state = "Non valide"
    else:
        license_state = "Non renseignée"

    table = document.add_table(rows=7, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_cell_text(table.cell(0, 0), "\nLIBELLE\n", bold=True)
    _set_cell_text(table.cell(0, 1), "\nVALEUR", bold=True)
    for cell in table.rows[0].cells:
        _shade(cell, "232323")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(0xFE, 0xD2, 0xF2)
    values = (
        ("\nHostname\n", _display(identity.hostname if identity else None)),
        ("\nVersion firmware\n", _display(identity.firmware_version if identity else None)),
        ("\nModèle de FortiGate\n", _display(identity.model if identity else None)),
        (
            "\nNuméro de série\n",
            _display((identity.serial_number if identity else None) or context.serial_number),
        ),
        ("\nCluster\n", cluster),
        ("\nLicences\n", license_state),
    )
    for row, (label, value) in enumerate(values, start=1):
        _set_cell_text(table.cell(row, 0), label)
        _set_cell_text(table.cell(row, 1), f"\n{value}\n")
        table.cell(row, 0).width = Cm(4)
        table.cell(row, 1).width = Cm(8)
    document.add_paragraph("")


def _add_utm_matrix(document: Any) -> None:
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = (
        "\nProfil de sécurité\n",
        "\nFlux entrant (vers réseau interne)\n",
        "\nFlux sortant (vers Internet)\n",
    )
    for cell, value in zip(table.rows[0].cells, headers, strict=True):
        _set_cell_text(cell, value, bold=True)
        _shade(cell, "232323")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(0xFE, 0xD2, 0xF2)
    _repeat_table_header(table.rows[0])
    rows = (
        ("\nFiltre DNS\n", "\nNon\n", "\nOui\n"),
        ("\nFiltre Web\n", "\nNon\n", "\nOui\n"),
        ("\nAntivirus\n", "\nOui\n", "\nOui\n"),
        ("\nIPS\n", "\nOui\n", "\nOui\n"),
        ("\nApplication Control\n", "\nOui\n", "\nOui\n"),
    )
    for row_values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values, strict=True):
            _set_cell_text(cell, value)
    document.add_paragraph("")


def _add_risk_summary(document: Any, risks: list[dict[str, str]]) -> None:
    if risks:
        document.add_page_break()
    document.add_heading(" Tableau récapitulatif des risques", level=2)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ("ID", "POINT AUDITE ET DESCRIPTION DU RISQUE", "RISQUE", "CORRECTION")
    for cell, value in zip(table.rows[0].cells, headers, strict=True):
        _set_cell_text(cell, value, bold=True)
        _shade(cell, "232323")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(0xFE, 0xD2, 0xF2)
    _repeat_table_header(table.rows[0])
    if not risks:
        cells = table.add_row().cells
        cells[0].merge(cells[3])
        _set_cell_text(
            cells[0],
            "Aucune non-conformité nécessitant un tableau de risque n'a été identifiée.",
        )
        return
    for risk in risks:
        cells = table.add_row().cells
        _set_cell_text(cells[0], risk["id"])
        _set_cell_text(cells[1], f"{risk['point']} : {risk['description']}")
        _set_cell_text(cells[2], risk["risk"])
        _set_cell_text(cells[3], risk["correction"])
        _shade(cells[2], _RISK_COLORS[risk["risk"]])
        _shade(cells[3], _CORRECTION_COLORS[risk["correction"]])


def _set_update_fields(document: Any) -> None:
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def _disable_template_heading_numbering(document: Any) -> None:
    """Use the explicit V1 numbers exactly once in generated headings."""

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        paragraph_properties = style._element.get_or_add_pPr()
        numbering = paragraph_properties.find(qn("w:numPr"))
        if numbering is not None:
            paragraph_properties.remove(numbering)


def _iter_document_paragraphs(document: Any):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        for container in (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            for paragraph in container.paragraphs:
                yield paragraph


def _insert_client_logo(document: Any, logo_path: str | Path | None) -> None:
    if logo_path is None:
        return
    path = Path(logo_path)
    if not path.is_file():
        return
    for paragraph in _iter_document_paragraphs(document):
        if "LOGO CLIENT si existe" not in paragraph.text:
            continue
        paragraph.text = ""
        run = paragraph.add_run()
        run.add_picture(str(path), width=Cm(4.0))
        return


def _insert_v1_toc(document: Any) -> None:
    introduction = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.strip() == "Introduction"
        ),
        None,
    )
    if introduction is None:
        return

    def add_entry(number: str, title: str, style: str) -> None:
        paragraph = introduction.insert_paragraph_before(style=style)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = False
        paragraph.add_run(f"{number}\t{title}")

    # Each TOC row must be its own paragraph.  A single paragraph containing
    # line breaks inherits the template's justified alignment, which stretches
    # every word across the page in Word/LibreOffice.
    add_entry("3.", "Audit de configuration", "toc 1")
    for section in V1_DOCUMENT_SECTIONS:
        add_entry(section.number, section.title, "toc 2")
        for slot in section.slots:
            add_entry(slot.number, slot.title, "toc 3")


def _add_asset_image(document: Any, filename: str | None) -> None:
    if filename is None:
        return
    path = Path(__file__).with_name("assets") / filename
    if not path.is_file():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Cm(15.5))
    document.add_paragraph("")


def _aggregate_status(findings: tuple[AuditFinding, ...]) -> AuditStatus:
    statuses = {finding.status for finding in findings}
    if AuditStatus.FAIL in statuses:
        return AuditStatus.FAIL
    if AuditStatus.ERROR in statuses or AuditStatus.UNKNOWN in statuses:
        return AuditStatus.UNKNOWN
    if statuses <= {AuditStatus.NOT_APPLICABLE}:
        return AuditStatus.NOT_APPLICABLE
    return AuditStatus.PASS


def _slot_finding(
    report: JsonAuditReport,
    presentation: AuditPresentation,
    slot: V1DocumentSlot,
) -> AuditFinding | None:
    findings_by_id = {finding.control_id: finding for finding in report.findings}
    actual = tuple(
        findings_by_id[control_id]
        for control_id in slot.control_ids
        if control_id in findings_by_id
    )
    if actual:
        status = _aggregate_status(actual)
        messages = tuple(dict.fromkeys(item.message for item in actual if item.message))
        primary = actual[0]
        return primary.model_copy(
            update={
                "title": slot.title,
                "display_name": slot.title,
                "status": status,
                "applicability": (
                    Applicability.NOT_APPLICABLE
                    if status is AuditStatus.NOT_APPLICABLE
                    else Applicability.APPLICABLE
                    if status in {AuditStatus.PASS, AuditStatus.FAIL}
                    else Applicability.UNKNOWN
                ),
                "message": " ".join(messages) or primary.message,
                "risk": next((item.risk for item in actual if item.risk is not None), primary.risk),
                "recommendation": next(
                    (item.recommendation for item in actual if item.recommendation),
                    primary.recommendation,
                ),
                "remediation": next(
                    (item.remediation for item in actual if item.remediation),
                    primary.remediation,
                ),
            }
        )

    row = next(
        (
            item
            for item in presentation.business_rows
            if set(slot.control_ids).intersection(item.v2_control_ids)
        ),
        None,
    )
    if row is not None and row.status is not None:
        return AuditFinding(
            control_id=slot.control_ids[0] if slot.control_ids else f"DOC-{slot.number}",
            title=slot.title,
            display_name=slot.title,
            status=row.status,
            applicability=row.applicability or Applicability.UNKNOWN,
            message=row.result or "Le résultat de ce point n'est pas disponible.",
        )

    if slot.kind == "model":
        return AuditFinding(
            control_id="DOC-MODEL-001",
            title=slot.title,
            display_name=slot.title,
            status=AuditStatus.UNKNOWN,
            applicability=Applicability.UNKNOWN,
            message=(
                "Le modèle FortiGate a été identifié, mais son statut de support actif "
                "doit être vérifié auprès d'une source EOL versionnée."
            ),
        )
    return None


def _render_slot(
    document: Any,
    report: JsonAuditReport,
    presentation: AuditPresentation,
    slot: V1DocumentSlot,
    risks: list[dict[str, str]],
    risk_number: int,
) -> int:
    if slot.kind == "group":
        document.add_heading(f"{slot.number} {slot.title}", level=3)
        document.add_paragraph(
            "Les contrôles historiques associés sont présentés dans les points suivants."
        )
        return risk_number
    if slot.kind == "cross_reference":
        document.add_heading(f"{slot.number} {slot.title}", level=3)
        document.add_paragraph(
            slot.note or "Ce point est détaillé dans le contrôle métier correspondant."
        )
        return risk_number

    finding = _slot_finding(report, presentation, slot)
    if finding is None:
        return risk_number

    document.add_heading(f"{slot.number} {slot.title}", level=3)
    document.add_paragraph("")
    point = document.add_paragraph()
    point.add_run("Point audité :").bold = True
    document.add_paragraph(_point_audited(finding))
    document.add_paragraph("")
    result = document.add_paragraph()
    result_run = result.add_run("Résultat : ")
    result_run.bold = True
    status_label = {
        AuditStatus.PASS: "CONFORME",
        AuditStatus.FAIL: "NON CONFORME",
        AuditStatus.NOT_APPLICABLE: "NON APPLICABLE",
    }.get(finding.status, "À VÉRIFIER")
    status_run = result.add_run(status_label)
    _color_status(status_run, finding.status)
    document.add_paragraph(_result_explanation(finding))
    document.add_paragraph("")
    if finding.status is AuditStatus.FAIL:
        risks.append(_add_risk_table(document, risk_number, finding))
        risk_number += 1
    _add_asset_image(document, slot.image)
    return risk_number


def render_docx(
    report: JsonAuditReport,
    client_logo_path: str | Path | None = None,
) -> bytes:
    """Render V2 findings in the historical V1 client-document order."""

    if not _TEMPLATE.is_file():
        raise FileNotFoundError(f"V1 DOCX template not found: {_TEMPLATE}")
    document = Document(str(_TEMPLATE))
    _disable_template_heading_numbering(document)
    context = report.context
    client = _display(context.client, "NOM CLIENT")
    site = _display(context.site, "SITE À RENSEIGNER")
    model = _display(report.device_identity.model if report.device_identity else None, "FortiGate")
    _replace_everywhere(
        document,
        {
            "NOM CLIENT": client,
            "LG AUTOS INVEST": client,
            "SAUSHEIM / ILLSACH": site,
            "Fortigate 200F": f"FortiGate {model}",
        },
    )
    _insert_client_logo(document, client_logo_path)
    _insert_v1_toc(document)
    document.core_properties.title = "Audit de configuration FortiGate"
    document.core_properties.subject = "Rapport d'audit de configuration"
    document.core_properties.author = "SNS Security"
    document.core_properties.created = report.created_at
    document.core_properties.modified = report.created_at
    _set_update_fields(document)

    document.add_page_break()
    document.add_heading("3. Audit de configuration", level=1)
    document.add_paragraph("")
    characteristics = document.add_paragraph()
    characteristics_run = characteristics.add_run("Caractéristiques :")
    characteristics_run.bold = True
    characteristics_run.font.size = Pt(14)
    document.add_paragraph("")
    _add_characteristics(document, report)

    presentation = _presentation_for(report)
    risks: list[dict[str, str]] = []
    risk_number = 1
    for section in V1_DOCUMENT_SECTIONS:
        document.add_heading(f"{section.number} {section.title}", level=2)
        document.add_paragraph("")
        for slot in section.slots:
            risk_number = _render_slot(
                document,
                report,
                presentation,
                slot,
                risks,
                risk_number,
            )
        if section.number == "3.5":
            _add_utm_matrix(document)

    _add_risk_summary(document, risks)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
