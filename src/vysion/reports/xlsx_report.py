from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from vysion.audit.models import AffectedObject, EvidenceItem, RiskAssessment
from vysion.reports.json_report import JsonAuditReport
from vysion.reports.presentation import AuditPresentation, build_presentation


def spreadsheet_text(value: object) -> str:
    text = "Non renseigné" if value is None or value == "" else str(value)
    if text.startswith(("=", "+", "-", "@")):
        return f"'{text}"
    return text


def evidence_text(items: tuple[EvidenceItem, ...]) -> str:
    if not items:
        return "Non renseigné"
    return "\n".join(
        ": ".join(
            part
            for part in (
                item.section,
                item.entry,
                item.directive,
                " ".join(item.tokens) or None,
                f"ligne {item.line}" if item.line else None,
                item.certainty.value,
                "defaulted" if item.defaulted else "explicit",
            )
            if part
        )
        for item in items
    )


def objects_text(objects: tuple[AffectedObject, ...]) -> str:
    return ", ".join(f"{item.object_type}: {item.name}" for item in objects) or "Non renseigné"


def risk_text(risk: RiskAssessment | None) -> str:
    if risk is None:
        return "Non renseigné"
    return "\n".join(
        part
        for part in (
            risk.summary,
            f"Impact: {risk.impact}" if risk.impact else None,
            f"Probabilité: {risk.likelihood}" if risk.likelihood else None,
            f"Traitement: {risk.treatment}" if risk.treatment else None,
        )
        if part
    )


def _presentation_for(report: JsonAuditReport) -> AuditPresentation:
    return report.presentation or build_presentation(report.findings)


def render_xlsx(report: JsonAuditReport) -> bytes:
    presentation = _presentation_for(report)
    workbook = Workbook()
    summary = workbook.active
    if summary is None:
        raise RuntimeError("workbook has no active worksheet")
    summary.title = "Synthèse"
    summary_rows = (
        ("Identifiant", str(report.report_id)),
        ("Source", spreadsheet_text(report.source_name)),
        ("Créé le", report.created_at.isoformat()),
        ("Expire le", report.expires_at.isoformat()),
        ("FortiGuard", report.fortiguard.status.value),
        ("Détail FortiGuard", spreadsheet_text(report.fortiguard.detail)),
        ("Client", spreadsheet_text(report.context.client)),
        ("Site", spreadsheet_text(report.context.site)),
        (
            "WAN sélectionnées",
            spreadsheet_text(", ".join(report.context.selected_wans or ()) or None),
        ),
        ("HA", spreadsheet_text(report.context.ha)),
        ("MPLS", spreadsheet_text(report.context.mpls)),
        ("Licence UTM", spreadsheet_text(report.context.utm_license)),
        (
            "Provenance opérateur",
            spreadsheet_text(
                report.context.operator_provenance.source
                if report.context.operator_provenance
                else None
            ),
        ),
        (
            "Opérateur",
            spreadsheet_text(
                report.context.operator_provenance.operator
                if report.context.operator_provenance
                else None
            ),
        ),
        ("Points métier V1 comparables", presentation.business_control_count),
        ("Contrôles moteur V2", presentation.engine_control_count),
        ("Sous-contrôles issus des splits", presentation.split_extra_finding_count),
        ("Contrôles complémentaires V2", presentation.v2_only_control_count),
        *(
            ("Explication du périmètre", line)
            for line in presentation.explanation_lines
        ),
    )
    for row in summary_rows:
        summary.append(row)
    for cell in summary["A"]:
        cell.font = Font(bold=True)
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 60

    controls = workbook.create_sheet("Contrôles")
    legacy_headers = (
        "Contrôle V2 (interne)",
        "Libellé métier V1",
        "Statut",
        "Constat",
        "Risque",
        "Recommandation",
    )
    controls.append(legacy_headers)
    for cell in controls[1]:
        cell.font = Font(bold=True)
    controls.freeze_panes = "A2"
    controls.auto_filter.ref = "A1:F1"
    for finding in report.findings:
        controls.append(
            (
                spreadsheet_text(finding.control_id),
                spreadsheet_text(finding.display_name or finding.title),
                spreadsheet_text(finding.status.value),
                spreadsheet_text(finding.message),
                spreadsheet_text(risk_text(finding.risk)),
                spreadsheet_text(finding.recommendation),
            )
        )
    for index, width in enumerate((22, 32, 12, 48, 60, 48), start=1):
        controls.column_dimensions[get_column_letter(index)].width = width

    enriched = workbook.create_sheet("Contrôles enrichis")
    headers = (
        "Contrôle V2 (interne)",
        "Libellé métier V1",
        "Catégorie",
        "Priorité",
        "Sévérité",
        "Applicabilité",
        "Statut",
        "Constat",
        "Preuve",
        "Objets affectés",
        "Risque",
        "Recommandation",
        "Remédiation",
        "Approbation client",
    )
    enriched.append(headers)
    for cell in enriched[1]:
        cell.font = Font(bold=True)
    enriched.freeze_panes = "A2"
    enriched.auto_filter.ref = "A1:N1"
    for finding in report.findings:
        enriched.append(
            tuple(
                spreadsheet_text(value)
                for value in (
                    finding.control_id,
                    finding.display_name or finding.title,
                    finding.category,
                    finding.priority.value,
                    finding.severity.value,
                    finding.applicability.value,
                    finding.status.value,
                    finding.message,
                    evidence_text(finding.evidence_items),
                    objects_text(finding.affected_objects),
                    risk_text(finding.risk),
                    finding.recommendation,
                    finding.remediation,
                    finding.customer_approval,
                )
            )
        )
    for index, width in enumerate(
        (22, 32, 18, 12, 14, 18, 12, 48, 70, 40, 70, 48, 60, 20),
        start=1,
    ):
        enriched.column_dimensions[get_column_letter(index)].width = width

    matrix = workbook.create_sheet("Matrice V1-V2")
    matrix_headers = (
        "Contrôle V1",
        "Libellé métier V1",
        "Contrôle(s) V2 correspondant(s)",
        "Relation",
        "Classification",
        "Résultat client",
    )
    matrix.append(matrix_headers)
    for cell in matrix[1]:
        cell.font = Font(bold=True)
    for row in (*presentation.business_rows, *presentation.v2_only_rows):
        targets = ", ".join(row.v2_control_ids) or (row.v2_projection or "Aucun finding moteur")
        matrix.append(
            (
                spreadsheet_text(row.business_key),
                spreadsheet_text(row.display_name),
                spreadsheet_text(targets),
                spreadsheet_text(row.relation),
                spreadsheet_text(row.classification),
                spreadsheet_text(
                    row.status.value if row.status is not None else row.result
                ),
            )
        )
    matrix.freeze_panes = "A2"
    matrix.auto_filter.ref = f"A1:F{matrix.max_row}"
    for index, width in enumerate((16, 64, 58, 22, 28, 64), start=1):
        matrix.column_dimensions[get_column_letter(index)].width = width

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
